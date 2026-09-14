from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "MPRIMS_Five_Day_Angular_Scaffolding_Plan.docx"

# Compact reference-guide palette: restrained enough for a technical working plan.
NAVY = "123B2A"
GREEN = "176B3A"
LIGHT_GREEN = "E8F2EC"
PALE = "F4F6F9"
GOLD = "D9A514"
INK = "17221C"
MUTED = "5F6F66"
WHITE = "FFFFFF"
RED = "9B1C1C"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent=120):
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[i]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_run(run, size=11, bold=False, color=INK, italic=False, font="Aptos"):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", *, size=11, bold=False, color=INK, italic=False,
             after=6, before=0, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    set_run(p.add_run(text), size=10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(text), size={1: 16, 2: 13, 3: 11.5}[level], bold=True,
            color=GREEN if level < 3 else NAVY)
    return p


def add_callout(doc, label, text, fill=LIGHT_GREEN, accent=GREEN):
    # A shaded paragraph is semantically cleaner than using a one-cell table as a box.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    set_run(p.add_run(label + "  "), size=10.5, bold=True, color=accent)
    set_run(p.add_run(text), size=10.5, color=INK)


def add_check(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("☐  " + text), size=10.5)


def add_day(doc, number, title, objective, tasks, deliverables, acceptance):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    set_run(p.add_run(f"DAY {number}"), size=10, bold=True, color=GOLD)
    set_run(p.add_run(f"  {title}"), size=16, bold=True, color=NAVY)
    add_callout(doc, "Objective", objective)
    add_heading(doc, "Work to complete", 3)
    for item in tasks:
        add_bullet(doc, item)
    add_heading(doc, "Deliverables", 3)
    for item in deliverables:
        add_bullet(doc, item)
    add_heading(doc, "Acceptance checks", 3)
    for item in acceptance:
        add_check(doc, item)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1)
section.right_margin = Inches(1)
section.header_distance = Inches(0.42)
section.footer_distance = Inches(0.42)

# Explicit style tokens for the compact_reference_guide preset.
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15
for level, size, before, after, color in ((1, 16, 18, 10, GREEN), (2, 13, 14, 7, GREEN), (3, 11.5, 10, 5, NAVY)):
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Quiet running furniture keeps the plan identifiable when printed.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run(hp.add_run("MPRIMS  |  FRONT-END SCAFFOLDING PLAN"), size=8.5, bold=True, color=MUTED)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_run(fp.add_run("Internal working plan  •  25 August 2026"), size=8.5, color=MUTED)

# First-page masthead.
add_para(doc, "MPRIMS DELIVERY PLAN", size=10, bold=True, color=GOLD, after=8)
add_para(doc, "Five-Day Angular + DevExtreme Scaffolding Plan", size=25, bold=True, color=NAVY, after=7)
add_para(doc, "An incremental front-end foundation for the Mpumalanga Provincial Road Incident Management System",
         size=12.5, color=MUTED, after=16)

meta = doc.add_table(rows=1, cols=2)
meta.style = "Table Grid"
set_table_geometry(meta, [2700, 6660])
for i, text in enumerate(("Plan detail", "Value")):
    set_cell_shading(meta.rows[0].cells[i], NAVY)
    set_run(meta.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
set_repeat_table_header(meta.rows[0])
for label, value in [
    ("Duration", "Five focused working days"),
    ("Technical direction", "Angular + DevExtreme UI, C# ASP.NET Core APIs, existing EF Core/SQL Server services"),
    ("Pilot feature", "Excel import → staged validation → issue review → approval → commit"),
]:
    row = meta.add_row()
    set_cell_shading(row.cells[0], PALE)
    set_run(row.cells[0].paragraphs[0].add_run(label), size=9.5, bold=True, color=NAVY)
    set_run(row.cells[1].paragraphs[0].add_run(value), size=9.5)

add_callout(doc, "End-of-week outcome", "A runnable Angular shell with DevExtreme, a secure C# API boundary, and one complete import-review vertical slice. The current Razor application remains available while features are migrated safely.", fill="FFF8E3", accent="7A5A00")

add_heading(doc, "1. Purpose and guardrails", 1)
add_para(doc, "The goal is to scaffold the future application structure and prove it against a real MPRIMS workflow. It is not realistic or necessary to rewrite the entire registry in five days.")
for item in [
    "Reuse the existing business services, validation rules, Identity roles/privileges, EF Core model, and SQL Server database.",
    "Place Angular and DevExtreme in front of C# API controllers; do not move registry rules into browser code.",
    "Migrate incrementally. Existing Razor pages remain the fallback until a replacement feature passes acceptance checks.",
    "Use the cleaned staging data as the reporting source only after review and approval; preserve audit history throughout.",
    "Confirm the organisation’s DevExpress/DevExtreme licence before production use.",
]:
    add_bullet(doc, item)

add_heading(doc, "2. Target architecture", 1)
for label, text in [
    ("Presentation", "Angular routes, reusable page shell, DevExtreme grids/forms/charts, responsive states, loading and error feedback."),
    ("API boundary", "C# ASP.NET Core API controllers returning DTOs and ProblemDetails responses; no EF entities exposed directly."),
    ("Application services", "Existing import, cleaning, validation, review, commit, report and authorisation services."),
    ("Data", "EF Core with SQL Server; staged rows remain separate from committed crash-registry records."),
    ("Security", "ASP.NET Identity, role/privilege policies, secure cookies for a same-host deployment, antiforgery protection and server-side validation."),
]:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(label + ": "), size=10.5, bold=True, color=GREEN)
    set_run(p.add_run(text), size=10.5)

add_heading(doc, "3. Five-day delivery map", 1)
overview = doc.add_table(rows=1, cols=4)
overview.style = "Table Grid"
set_table_geometry(overview, [900, 2100, 3300, 3060])
headers = ["Day", "Theme", "Primary output", "Proof of completion"]
for i, text in enumerate(headers):
    set_cell_shading(overview.rows[0].cells[i], NAVY)
    set_run(overview.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
set_repeat_table_header(overview.rows[0])
rows = [
    ("1", "Foundation", "Angular/DevExtreme workspace and shell", "Build runs; routes and layouts render"),
    ("2", "API boundary", "Secure import API contracts", "Swagger/API tests return stable DTOs"),
    ("3", "Angular services", "Typed client and upload workflow", "Workbook reaches staging through UI"),
    ("4", "Review slice", "Issue review, decisions and commit UI", "Happy path and blocking path work"),
    ("5", "Hardening", "Tested handover and migration backlog", "Demo, security checks and docs pass"),
]
for values in rows:
    cells = overview.add_row().cells
    for i, value in enumerate(values):
        set_run(cells[i].paragraphs[0].add_run(value), size=8.8, bold=(i == 0), color=NAVY if i == 0 else INK)
        if int(values[0]) % 2 == 0:
            set_cell_shading(cells[i], "F7FAF8")

doc.add_page_break()
add_heading(doc, "4. Daily execution plan", 1)

add_day(doc, 1, "Foundation and technical baseline",
        "Create a clean front-end foundation without disturbing the working MVC application.",
        [
            "Record the current build, routes, authentication behaviour, import workflow and known DevExtreme licence warnings.",
            "Create the Angular workspace in a separate client folder/solution area and add DevExtreme using the approved package version.",
            "Choose same-host deployment for the first scaffold so ASP.NET Identity cookies can be reused with less CORS/token complexity.",
            "Create the application shell: header, side navigation, route outlet, access-denied page, not-found page and responsive breakpoints.",
            "Define feature folders for dashboard, imports, crash records, reports, insights and administration.",
            "Add environment configuration and a development proxy to the ASP.NET Core host.",
        ],
        ["Angular application builds and starts.", "DevExtreme theme and one test component render.", "Architecture/readme note documents the folder structure and local run commands."],
        ["Existing .NET solution still builds.", "Navigation works at desktop and narrow widths.", "No secret, connection string or user data is committed to the client."],)

add_day(doc, 2, "C# API boundary and security",
        "Expose stable, documented endpoints around existing services while keeping business rules on the server.",
        [
            "Create C# ControllerBase API controllers for import batches, upload, staged summaries, quality issues, row decisions, approval and commit.",
            "Introduce request/response DTOs and mapping; do not serialize EF Core entities or navigation properties directly.",
            "Apply existing privilege policies to every endpoint, including distinct permissions for upload, review, approval and commit.",
            "Return consistent validation and failure responses using ProblemDetails; include a correlation ID for support diagnostics.",
            "Configure secure cookies, antiforgery handling, upload limits, allowed file extensions, cancellation tokens and server-side validation.",
            "Add endpoint-level tests for authenticated, forbidden, invalid and successful requests.",
        ],
        ["Versioned or clearly grouped import endpoints.", "DTOs and commented controller code.", "API test collection or automated integration tests."],
        ["Anonymous and under-privileged requests are rejected.", "Invalid workbooks produce useful 4xx responses rather than 500 errors.", "No database entity graph leaks through JSON."],)

doc.add_page_break()
add_day(doc, 3, "Angular data layer and workbook intake",
        "Connect the new UI to the C# APIs with typed, reusable client-side infrastructure.",
        [
            "Configure HttpClient, API base URL, credentials, antiforgery header handling and a correlation-ID/error interceptor.",
            "Create TypeScript interfaces that mirror API DTOs, plus services for import batches, staged rows and issue decisions.",
            "Build the import landing page with a DevExtreme uploader, progress state, workbook requirements and clear validation feedback.",
            "Add batch history with status, uploader, dates, row totals, warning/blocking counts and actions.",
            "Handle loading, empty, offline, forbidden and unexpected-error states consistently.",
            "Write focused component/service tests for upload and batch loading.",
        ],
        ["Typed API client and interceptors.", "Workbook upload page and batch history grid.", "Reusable notification/loading patterns."],
        ["A valid workbook creates a staging batch.", "An invalid file explains what the user must fix.", "Refresh preserves server truth; the browser is not the source of record."],)

add_day(doc, 4, "Data-quality review vertical slice",
        "Deliver the first complete business workflow in Angular: inspect, resolve, approve and safely commit cleaned data.",
        [
            "Build a DevExtreme review grid with server-side paging/filtering, row status, original values, suggested values and severity badges.",
            "Add master-detail or a side panel for field-level issues, source context, resolution notes and audit history.",
            "Implement actions for accept suggestion, edit value, reject row, approve row, batch approval and commit, subject to privileges.",
            "Require confirmation for commit and other irreversible workflow actions; display blocking issue counts before the user proceeds.",
            "Prevent double submission and handle concurrency conflicts with a refresh-and-retry message.",
            "Verify that committed records originate from approved staging rows and that issue decisions remain auditable.",
        ],
        ["Review workspace with realistic data states.", "Role-aware actions and confirmations.", "Successful end-to-end upload-to-commit demonstration."],
        ["Blocking issues stop approval/commit.", "Resolved values are visible after reload.", "Audit fields show who acted and when.", "Keyboard navigation and labels work for core actions."],)

doc.add_page_break()
add_day(doc, 5, "Hardening, demonstration and handover",
        "Leave a reliable scaffold the team can extend, not a one-off demonstration.",
        [
            "Run the .NET build, Angular production build, API tests, component tests and a complete workbook smoke test.",
            "Review security: authorisation coverage, antiforgery, upload validation, sensitive logging, dependency warnings and client-side secrets.",
            "Review UX on common desktop sizes and a narrow viewport; fix overflow, inaccessible labels and unclear empty/error states.",
            "Document local setup, run commands, API conventions, feature-folder pattern and how Razor and Angular coexist during migration.",
            "Create the next-feature backlog: dashboard, crash registry, centralized reports, insights, lookup administration and user/role management.",
            "Demonstrate the vertical slice to a senior/stakeholder and record accepted gaps and decisions.",
        ],
        ["Green build/test evidence.", "Developer handover guide and known-issues list.", "Prioritised migration backlog with acceptance criteria."],
        ["A new developer can run both applications from the instructions.", "The import workflow works with a representative workbook.", "Razor fallback remains available.", "DevExtreme licensing and deployment ownership are recorded."],)

add_heading(doc, "5. Security and UX baseline", 1)
security = doc.add_table(rows=1, cols=3)
security.style = "Table Grid"
set_table_geometry(security, [2100, 3900, 3360])
for i, text in enumerate(["Area", "Required scaffold", "Five-day evidence"]):
    set_cell_shading(security.rows[0].cells[i], NAVY)
    set_run(security.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
set_repeat_table_header(security.rows[0])
for area, req, evidence in [
    ("Authentication", "ASP.NET Identity; same-host secure cookies", "Unauthenticated API test fails"),
    ("Authorisation", "Privilege policy on UI and API; API is authoritative", "Forbidden tests for upload/review/commit"),
    ("Request safety", "Antiforgery, size/type limits, server validation", "Invalid and forged requests rejected"),
    ("Data exposure", "DTOs, safe logs, no client secrets", "Response and logging review"),
    ("Auditability", "Actor, time, decision and notes retained", "Reloaded review history"),
    ("UX states", "Loading, empty, success, warning, error, forbidden", "State checklist on pilot pages"),
    ("Accessibility", "Labels, keyboard flow, focus, contrast", "Manual keyboard/contrast check"),
]:
    cells = security.add_row().cells
    for i, value in enumerate((area, req, evidence)):
        set_run(cells[i].paragraphs[0].add_run(value), size=8.7, bold=(i == 0), color=NAVY if i == 0 else INK)

add_heading(doc, "6. Roles and working rhythm", 1)
add_para(doc, "If one developer owns the scaffold, use the first 20 minutes of each morning to confirm the day’s acceptance checks and the final 30 minutes to build, test and document. If more people are available, split front-end, API and review responsibilities but integrate at least twice daily.")
for item in [
    "Developer: implements the slice, writes comments around non-obvious decisions, and keeps builds green.",
    "Product/data owner: supplies a representative workbook and confirms validation, terminology and review decisions.",
    "Security/technical reviewer: checks endpoint policies, deployment assumptions and dependency/licensing decisions.",
    "Senior stakeholder: attends the Day 5 demonstration and approves the next migration slice.",
]:
    add_bullet(doc, item)

add_heading(doc, "7. Key risks and responses", 1)
risks = [
    ("Scope expands into a full rewrite", "Keep the five-day goal to shell + API boundary + one vertical slice."),
    ("Angular duplicates C# rules", "Treat client validation as guidance only; server services remain authoritative."),
    ("Authentication becomes complex", "Start same-host with Identity cookies; revisit tokens only for a real deployment need."),
    ("Large imports make grids slow", "Use server-side paging/filtering and return summary counts separately."),
    ("Old and new UI diverge", "Use one set of services and DTO-backed APIs; migrate feature by feature with acceptance gates."),
    ("DevExtreme licence is unresolved", "Confirm entitlement before production deployment and record the decision."),
]
for risk, response in risks:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    set_run(p.add_run(risk + " — "), size=10.3, bold=True, color=RED)
    set_run(p.add_run(response), size=10.3)

add_heading(doc, "8. Definition of done", 1)
for item in [
    "Both the existing .NET application and the new Angular client build from documented commands.",
    "Angular has a reusable shell, route structure, DevExtreme configuration and consistent application states.",
    "C# APIs are DTO-based, validated, authorised, tested and connected to existing business services.",
    "A user can upload a workbook, review staged issues, resolve/approve rows and commit eligible data end to end.",
    "The workflow preserves staging separation, blocking rules and audit history.",
    "Security, accessibility, responsive behaviour, licensing and known gaps are documented.",
    "The current Razor feature remains available until the new slice is accepted for cutover.",
    "The next migration backlog is prioritised and ready for estimation.",
]:
    add_check(doc, item)

add_callout(doc, "Recommended next slice", "After the import-review pilot is accepted, scaffold the centralized Reports area because it exercises shared filters, permissions, asynchronous generation, download history and the monthly/quarterly/6-month/9-month report family without changing the core crash-capture workflow.", fill="FFF8E3", accent="7A5A00")

# Core document metadata.
doc.core_properties.title = "MPRIMS Five-Day Angular and DevExtreme Scaffolding Plan"
doc.core_properties.subject = "Incremental front-end scaffolding and import-review vertical slice"
doc.core_properties.author = "MPRIMS Project Team"
doc.core_properties.keywords = "MPRIMS, Angular, DevExtreme, ASP.NET Core, scaffolding, data cleaning"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
