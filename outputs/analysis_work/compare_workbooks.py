import json
import math
import re
import statistics
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


FILES = {
    "original": Path(r"C:\Users\ReaobakaNtoagae\OneDrive - atnetgroup.com\EHL MARCH 2026.xlsx"),
    "converted": Path(r"C:\Users\ReaobakaNtoagae\Downloads\Upload EHLANZENI 2026-03.xlsx"),
    "filled": Path(r"C:\Users\ReaobakaNtoagae\Downloads\Upload EHLANZENI 2026-03 (filled).xlsx"),
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = re.sub(r"\s+", " ", value.strip())
        return value or None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def json_value(value):
    value = clean(value)
    if isinstance(value, (int, float, bool)) or value is None:
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return str(value)
        return value
    return str(value)


def row_values(ws, row_no, max_col=None):
    max_col = max_col or ws.max_column
    return [clean(ws.cell(row_no, col).value) for col in range(1, max_col + 1)]


def find_dense_header(ws):
    # Prefer rows with many text labels and a high proportion of non-empty cells.
    candidates = []
    for r in range(1, min(ws.max_row, 40) + 1):
        vals = row_values(ws, r)
        nonempty = [v for v in vals if v is not None]
        text_count = sum(isinstance(v, str) for v in nonempty)
        if nonempty:
            candidates.append((text_count, len(nonempty), -r, r))
    return max(candidates)[-1] if candidates else None


def merged_ranges(ws):
    try:
        return [str(x) for x in ws.merged_cells.ranges]
    except Exception:
        return []


def sheet_profile(ws_formula, ws_values):
    nonempty = []
    formulas = []
    values_by_row = defaultdict(list)
    styles = Counter()
    for row in ws_formula.iter_rows():
        for c in row:
            v = clean(c.value)
            if v is not None:
                nonempty.append((c.coordinate, v))
                values_by_row[c.row].append(v)
                styles[c.style_id] += 1
                if isinstance(c.value, str) and c.value.startswith("="):
                    formulas.append((c.coordinate, c.value, ws_values[c.coordinate].value))
    header_row = find_dense_header(ws_values)
    header = row_values(ws_values, header_row) if header_row else []
    return {
        "title": ws_formula.title,
        "visibility": ws_formula.sheet_state,
        "dimensions": ws_formula.calculate_dimension(),
        "max_row": ws_formula.max_row,
        "max_column": ws_formula.max_column,
        "nonempty_cells": len(nonempty),
        "nonempty_rows": len(values_by_row),
        "formula_count": len(formulas),
        "formulas_sample": [[a, b, json_value(c)] for a, b, c in formulas[:15]],
        "merged_ranges": merged_ranges(ws_formula),
        "style_ids": dict(styles),
        "freeze_panes": str(ws_formula.freeze_panes) if ws_formula.freeze_panes else None,
        "auto_filter": str(ws_formula.auto_filter.ref) if ws_formula.auto_filter.ref else None,
        "tables": list(ws_formula.tables.keys()),
        "data_validations": len(ws_formula.data_validations.dataValidation),
        "header_candidate_row": header_row,
        "header_candidate": [json_value(v) for v in header],
        "sample_nonempty": [[a, json_value(v)] for a, v in nonempty[:40]],
        "row_samples": {
            str(r): [json_value(v) for v in values_by_row[r]]
            for r in sorted(values_by_row)[:12]
        },
    }


def workbook_profile(path):
    # Read formula and cached-value views so we can distinguish calculations from hardcoded output.
    wb_formula = load_workbook(path, read_only=False, data_only=False, keep_links=True)
    wb_values = load_workbook(path, read_only=False, data_only=True, keep_links=True)
    profile = {
        "path": str(path),
        "sheet_names": wb_formula.sheetnames,
        "defined_names": sorted(str(n) for n in wb_formula.defined_names),
        "calculation": {
            "mode": getattr(wb_formula.calculation, "calcMode", None),
            "full_calc_on_load": getattr(wb_formula.calculation, "fullCalcOnLoad", None),
            "force_full_calc": getattr(wb_formula.calculation, "forceFullCalc", None),
        },
        "sheets": [],
    }
    for name in wb_formula.sheetnames:
        profile["sheets"].append(sheet_profile(wb_formula[name], wb_values[name]))
    return profile


def rectangular_records(ws, header_row):
    headers = []
    for col in range(1, ws.max_column + 1):
        value = clean(ws.cell(header_row, col).value)
        headers.append(str(value) if value is not None else f"__blank_{col}")
    records = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = [clean(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in values):
            records.append({headers[i]: json_value(v) for i, v in enumerate(values)})
    return headers, records


def normalized_text_counts(profile):
    counts = Counter()
    for sheet in profile["sheets"]:
        for _, value in sheet["sample_nonempty"]:
            if isinstance(value, str):
                counts[value.casefold()] += 1
    return counts


def main():
    result = {"profiles": {}, "tabular": {}, "errors": {}}
    workbooks = {}
    for key, path in FILES.items():
        try:
            result["profiles"][key] = workbook_profile(path)
            workbooks[key] = load_workbook(path, read_only=False, data_only=True)
        except Exception as exc:
            result["errors"][key] = repr(exc)

    for key, wb in workbooks.items():
        result["tabular"][key] = []
        for ws in wb.worksheets:
            header_row = find_dense_header(ws)
            if header_row:
                headers, records = rectangular_records(ws, header_row)
                result["tabular"][key].append({
                    "sheet": ws.title,
                    "header_row": header_row,
                    "headers": headers,
                    "record_count": len(records),
                    "records_sample": records[:8],
                    "column_nonblank": {
                        h: sum(1 for rec in records if rec[h] is not None)
                        for h in headers
                    },
                    "column_unique": {
                        h: len({str(rec[h]).casefold() for rec in records if rec[h] is not None})
                        for h in headers
                    },
                })

    out = Path(__file__).with_name("workbook_profiles.json")
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
