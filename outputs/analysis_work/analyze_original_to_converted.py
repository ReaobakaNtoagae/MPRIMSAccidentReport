import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


ORIGINAL = Path(r"C:\Users\ReaobakaNtoagae\OneDrive - atnetgroup.com\EHL MARCH 2026.xlsx")
CONVERTED = Path(r"C:\Users\ReaobakaNtoagae\Downloads\Upload EHLANZENI 2026-03.xlsx")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).strip()).upper()


def main():
    owb = load_workbook(ORIGINAL, data_only=True)
    cwb = load_workbook(CONVERTED, data_only=True)
    src = owb["MARCH 2026 CAS"]
    dst = cwb["MARCH 2026"]

    # These are the columns retained when the original 24-column structure became 21 columns.
    mapping = {
        "A": "A", "B": "B", "D": "C", "E": "D", "F": "E", "G": "F", "H": "G", "I": "H",
        "J": "I", "K": "J", "L": "K", "M": "L",
        "P": "M", "Q": "N", "R": "O", "S": "P",
        "T": "Q", "U": "R", "V": "S", "W": "T", "X": "U",
    }

    def src_key(r):
        # Exclude station so spelling normalisation does not prevent a match.
        return tuple(norm(src[f"{c}{r}"].value) for c in ("B", "D", "F", "H"))

    def dst_key(r):
        return tuple(norm(dst[f"{c}{r}"].value) for c in ("B", "C", "E", "G"))

    source_by_key = defaultdict(list)
    dest_by_key = defaultdict(list)
    for r in range(7, 236):
        source_by_key[src_key(r)].append(r)
        dest_by_key[dst_key(r)].append(r)

    matched = []
    unmatched_source = []
    unmatched_dest = []
    ambiguous_keys = []
    for key in sorted(set(source_by_key) | set(dest_by_key)):
        sr = source_by_key.get(key, [])
        dr = dest_by_key.get(key, [])
        if len(sr) == 1 and len(dr) == 1:
            matched.append((sr[0], dr[0]))
        elif len(sr) == len(dr) and len(sr) > 1:
            # Resolve repeated keys deterministically by the complete retained fingerprint.
            s_remaining = sr[:]
            d_remaining = dr[:]
            for srow in sr:
                sfp = tuple(norm(src[f"{c}{srow}"].value) for c in mapping)
                exact = [drow for drow in d_remaining if tuple(norm(dst[f"{mapping[c]}{drow}"].value) for c in mapping) == sfp]
                if exact:
                    drow = exact[0]
                    matched.append((srow, drow))
                    s_remaining.remove(srow)
                    d_remaining.remove(drow)
            for srow, drow in zip(s_remaining, d_remaining):
                matched.append((srow, drow))
            ambiguous_keys.append({"key": key, "source_rows": sr, "dest_rows": dr})
        else:
            unmatched_source.extend(sr)
            unmatched_dest.extend(dr)

    field_differences = Counter()
    difference_examples = defaultdict(list)
    station_changes = Counter()
    for sr, dr in matched:
        for sc, dc in mapping.items():
            sv = norm(src[f"{sc}{sr}"].value)
            dv = norm(dst[f"{dc}{dr}"].value)
            if sv != dv:
                field_differences[(sc, dc)] += 1
                if len(difference_examples[(sc, dc)]) < 12:
                    difference_examples[(sc, dc)].append({"source_row": sr, "dest_row": dr, "source": src[f"{sc}{sr}"].value, "dest": dst[f"{dc}{dr}"].value})
                if sc == "A":
                    station_changes[(sv, dv)] += 1

    cas_values = [norm(src[f"C{r}"].value) for r in range(7, 236)]
    male_values = [norm(src[f"N{r}"].value) for r in range(7, 236)]
    female_values = [norm(src[f"O{r}"].value) for r in range(7, 236)]

    original_preserved_rows = [tuple(norm(src[f"{c}{r}"].value) for c in mapping) for r in range(7, 236)]
    converted_rows = [tuple(norm(dst[f"{mapping[c]}{r}"].value) for c in mapping) for r in range(7, 236)]
    source_counter, dest_counter = Counter(original_preserved_rows), Counter(converted_rows)
    exact_multiset_overlap = sum((source_counter & dest_counter).values())

    result = {
        "source_records": 229,
        "converted_records": 229,
        "unique_key_matches": len(matched),
        "unmatched_source_rows": unmatched_source,
        "unmatched_destination_rows": unmatched_dest,
        "ambiguous_match_groups": ambiguous_keys,
        "exact_preserved_row_overlap": exact_multiset_overlap,
        "field_differences": [
            {"source_col": k[0], "dest_col": k[1], "count": v, "examples": difference_examples[k]}
            for k, v in field_differences.items()
        ],
        "station_changes": [
            {"source": k[0], "dest": k[1], "count": v} for k, v in station_changes.items()
        ],
        "dropped_fields": {
            "CAS_nonblank": sum(bool(v) for v in cas_values),
            "CAS_unique": len(set(v for v in cas_values if v)),
            "male_nonblank": sum(bool(v) for v in male_values),
            "female_nonblank": sum(bool(v) for v in female_values),
            "male_total": sum(float(v) for v in male_values if v.replace('.', '', 1).isdigit()),
            "female_total": sum(float(v) for v in female_values if v.replace('.', '', 1).isdigit()),
        },
        "original_supplementary_sections": {
            "victim_age_rows": "A240:F244",
            "victim_gender_rows": "A246:C252",
            "race_rows_begin": "A254:F254",
            "nonempty_rows_after_main_table": sum(1 for r in range(240, src.max_row + 1) if any(src.cell(r, c).value is not None for c in range(1, src.max_column + 1))),
        },
    }

    out = Path(__file__).with_name("original_to_converted_analysis.json")
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False, default=list), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
