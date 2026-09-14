import json
import re
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


CONVERTED = Path(r"C:\Users\ReaobakaNtoagae\Downloads\Upload EHLANZENI 2026-03.xlsx")
FILLED = Path(r"C:\Users\ReaobakaNtoagae\Downloads\Upload EHLANZENI 2026-03 (filled).xlsx")


def norm(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v).strip()).upper()


def main():
    source_wb = load_workbook(CONVERTED, data_only=True)
    target_wb = load_workbook(FILLED, data_only=True)
    source = source_wb["MARCH 2026"]
    target = target_wb["UPLOAD TEMPLATE"]

    # Source-to-template positions, including the three intentional new metadata fields.
    mapping = {
        "A": "A", "B": "B", "C": "D", "D": "E", "E": "F", "F": "G", "G": "H", "H": "I",
        "I": "J", "J": "K", "K": "L", "L": "M",
        "M": "Q", "N": "R", "O": "S", "P": "T",
        "Q": "U", "R": "V", "S": "W", "T": "X", "U": "Y",
    }

    mismatches = []
    raw_type_changes = Counter()
    whitespace_cleaned = Counter()
    for row in range(7, 236):
        for src_col, dst_col in mapping.items():
            src = source[f"{src_col}{row}"].value
            dst = target[f"{dst_col}{row}"].value
            if norm(src) != norm(dst):
                mismatches.append({"row": row, "source_cell": f"{src_col}{row}", "target_cell": f"{dst_col}{row}", "source": src, "target": dst})
            if src is not None and dst is not None and type(src).__name__ != type(dst).__name__:
                raw_type_changes[(src_col, dst_col, type(src).__name__, type(dst).__name__)] += 1
            if isinstance(src, str) and isinstance(dst, str) and src != dst and norm(src) == norm(dst):
                whitespace_cleaned[(src_col, dst_col)] += 1

    source_rows = [tuple(norm(source.cell(r, c).value) for c in range(1, 22)) for r in range(7, 236)]
    exact_dups = defaultdict(list)
    for row_num, values in zip(range(7, 236), source_rows):
        exact_dups[values].append(row_num)
    exact_dups = {"|".join(k): v for k, v in exact_dups.items() if len(v) > 1}

    # A practical business key for suspected duplicates; not a declaration that they are duplicates.
    key_dups = defaultdict(list)
    for r in range(7, 236):
        key = tuple(norm(source.cell(r, c).value) for c in (1, 2, 3, 5, 6, 7))
        key_dups[key].append(r)
    key_dups = {"|".join(k): v for k, v in key_dups.items() if len(v) > 1}

    station_list = {norm(source_wb["Stations"].cell(r, 1).value) for r in range(2, 27)}
    station_values = Counter(norm(source.cell(r, 1).value) for r in range(7, 236))
    outside = {k: v for k, v in station_values.items() if k not in station_list}
    unused = sorted(station_list - set(station_values))

    column_missing = {}
    for col in range(1, 30):
        header = target.cell(6, col).value
        key = f"{target.cell(6, col).column_letter}:{header}"
        column_missing[key] = sum(1 for r in range(7, 236) if norm(target.cell(r, col).value) == "")

    date_values = Counter(norm(source.cell(r, 3).value) for r in range(7, 236))
    day_values = Counter(norm(source.cell(r, 4).value) for r in range(7, 236))
    types = Counter(norm(source.cell(r, 8).value) for r in range(7, 236))
    routes = Counter(norm(source.cell(r, 6).value) for r in range(7, 236))

    result = {
        "row_count": 229,
        "mapped_cell_count": 229 * len(mapping),
        "normalized_mismatch_count": len(mismatches),
        "mismatches": mismatches,
        "raw_type_changes": [{"source_col": k[0], "target_col": k[1], "from": k[2], "to": k[3], "count": v} for k, v in raw_type_changes.items()],
        "whitespace_cleaned": [{"source_col": k[0], "target_col": k[1], "count": v} for k, v in whitespace_cleaned.items()],
        "exact_duplicate_groups": list(exact_dups.values()),
        "business_key_duplicate_groups": list(key_dups.values()),
        "stations_in_data": station_values,
        "stations_outside_reference": outside,
        "unused_reference_stations": unused,
        "column_missing_counts": column_missing,
        "date_values": date_values,
        "day_values": day_values,
        "accident_types": types,
        "routes": routes,
        "constant_metadata": {
            "region_values": Counter(norm(target[f"AA{r}"].value) for r in range(7, 236)),
            "year_values": Counter(norm(target[f"AB{r}"].value) for r in range(7, 236)),
            "month_values": Counter(norm(target[f"AC{r}"].value) for r in range(7, 236)),
        },
    }
    out = Path(__file__).with_name("conversion_analysis.json")
    out.write_text(json.dumps(result, indent=2, ensure_ascii=False, default=dict), encoding="utf-8")
    print(out)


if __name__ == "__main__":
    main()
