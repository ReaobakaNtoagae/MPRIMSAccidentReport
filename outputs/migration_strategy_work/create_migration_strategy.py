from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Accident_Reports_Rapid_Migration_Strategy.docx"

# Product-inspired palette: vivid cyan and aqua, balanced by formal dark blue.
NAVY = "082F49"
DEEP_BLUE = "0B3B5C"
CYAN = "06B6D4"
AQUA = "22D3EE"
PALE_CYAN = "ECFEFF"
PALE_BLUE = "F2F8FA"
WHITE = "FFFFFF"
INK = "142B38"
MUTED = "5F7480"
LINE = "CFE8ED"
GREEN = "0F8A6A"
GOLD = "B97800"
RED = "B63552"


def font(run, size=10.5, bold=False, color=INK, italic=False, family="Aptos"):
    run.font.name = family
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), family)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), family)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    node = tc_pr.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        tc_pr.append(node)
    node.set(qn("w:fill"), fill)


def cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_geometry(table, widths, indent=120):
    # Explicit DXA geometry prevents Word and LibreOffice from resizing tables differently.
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def paragraph(doc, text="", size=10.5, bold=False, color=INK, italic=False,
              before=0, after=6, align=None, keep=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    p.paragraph_format.keep_with_next = keep
    if align is not None:
        p.alignment = align
    font(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    font(p.add_run(text), size={1: 16, 2: 13, 3: 11.5}[level], bold=True,
         color=DEEP_BLUE if level == 1 else NAVY, family="Aptos Display")
    return p


def bullet(doc, text, level=0, checked=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    prefix = ""
    if checked is True:
        prefix = "☑  "
    elif checked is False:
        prefix = "☐  "
    font(p.add_run(prefix + text), size=10.2)
    return p


def callout(doc, label, text, fill=PALE_CYAN, accent=CYAN):
    # A shaded paragraph keeps the callout semantic and accessible without a layout table.
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.08)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(9)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    p_bdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "20")
    left.set(qn("w:space"), "7")
    left.set(qn("w:color"), accent)
    p_bdr.append(left)
    p_pr.append(p_bdr)
    font(p.add_run(label + "  "), size=10.4, bold=True, color=accent)
    font(p.add_run(text), size=10.4)


def phase(doc, number, title, effort, outcome, actions, gate):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.space_before = Pt(15)
    p.paragraph_format.space_after = Pt(5)
    font(p.add_run(f"PHASE {number}"), size=9.5, bold=True, color=CYAN)
    font(p.add_run(f"  {title}"), size=14, bold=True, color=NAVY, family="Aptos Display")
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(7)
    font(meta.add_run(f"Effort: {effort}  |  Outcome: "), size=9.5, bold=True, color=MUTED)
    font(meta.add_run(outcome), size=9.5, color=MUTED)
    for item in actions:
        bullet(doc, item)
    callout(doc, "Exit gate", gate, fill="E9F8F3", accent=GREEN)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.header_distance = Inches(0.38)
section.footer_distance = Inches(0.4)

# Resolve the compact-reference preset into explicit style values.
normal = doc.styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.18
for level, size, before, after, color in ((1, 16, 18, 9, DEEP_BLUE), (2, 13, 14, 7, NAVY), (3, 11.5, 10, 5, NAVY)):
    style = doc.styles[f"Heading {level}"]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Running header/footer use the new product identity without overwhelming the page.
hp = section.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(hp.add_run("ACCIDENT REPORTS ADMIN SYSTEM  /  MIGRATION STRATEGY"), size=8.2, bold=True, color=MUTED)
fp = section.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(fp.add_run("Rapid delivery • controlled migration • measurable quality"), size=8.2, color=MUTED)

# Customer-pack style cover.
paragraph(doc, "MIGRATION PLAYBOOK", size=10, bold=True, color=CYAN, after=5)
paragraph(doc, "From MVC to Angular + .NET Core API", size=26, bold=True, color=NAVY, after=5)
paragraph(doc, "A rapid, backend-led strategy for the Accident Reports Admin System",
          size=13, color=DEEP_BLUE, after=15)

meta = doc.add_table(rows=1, cols=2)
meta.style = "Table Grid"
table_geometry(meta, [2800, 6560])
for i, text in enumerate(("Plan detail", "Decision")):
    shade(meta.rows[0].cells[i], NAVY)
    font(meta.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
repeat_header(meta.rows[0])
for label, value in [
    ("Delivery model", "Backend-led vertical slices; Angular follows one stable feature behind"),
    ("Runtime structure", "AccidentReports.Api + AccidentReports.Core + Angular client"),
    ("Quality safety net", "AccidentReports.Tests plus repeatable workbook and API regression tests"),
    ("Database", "Retain the existing SQL Server database; one EF migration owner"),
    ("First production slice", "Workbook upload → staging → review → approval → commit"),
]:
    row = meta.add_row()
    shade(row.cells[0], PALE_BLUE)
    font(row.cells[0].paragraphs[0].add_run(label), size=9.2, bold=True, color=NAVY)
    font(row.cells[1].paragraphs[0].add_run(value), size=9.2)

callout(doc, "Mission", "Reach a usable, defensible new system quickly without rewriting proven business logic, losing data, or creating two competing sources of truth.")
callout(doc, "The honest timeline", "A credible import-focused MVP is approximately 10–15 focused working days after the current backend is stabilised. A responsible migration of the broader system is approximately 4–6 weeks of solo development, depending on acceptance feedback and the complexity of full crash capture.", fill="FFF6DD", accent=GOLD)

heading(doc, "1. Executive decision", 1)
paragraph(doc, "Use a strangler migration: the existing MVC application remains available while new API-backed Angular features replace it one complete workflow at a time. The backend leads by one feature, not by the entire system.")
for item in [
    "Do not perform a clean-sheet rewrite.",
    "Do not create a second production database.",
    "Do not migrate all controllers before proving one vertical slice.",
    "Do not put data-cleaning, approval or authorisation rules in Angular.",
    "Do not retire an MVC feature until its replacement passes acceptance testing.",
]:
    bullet(doc, item)

heading(doc, "2. What ‘fast with quality’ means", 1)
speed = doc.add_table(rows=1, cols=3)
speed.style = "Table Grid"
table_geometry(speed, [2500, 3430, 3430])
for i, text in enumerate(("Optimise", "Do", "Avoid")):
    shade(speed.rows[0].cells[i], NAVY)
    font(speed.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
repeat_header(speed.rows[0])
for a, b, c in [
    ("Decision speed", "Record a decision and continue", "Revisiting settled architecture daily"),
    ("Development", "Complete one end-to-end slice", "Starting six half-finished modules"),
    ("Testing", "Automate repeatable high-risk checks", "Depending on memory and manual retesting"),
    ("Database", "Use reviewed additive migrations", "Renaming or rebuilding the schema during UI migration"),
    ("Energy", "Use focused blocks and stop on fatigue", "Late-night production DB or security changes"),
]:
    row = speed.add_row()
    for i, value in enumerate((a, b, c)):
        font(row.cells[i].paragraphs[0].add_run(value), size=8.8, bold=(i == 0), color=NAVY if i == 0 else INK)

callout(doc, "Non-negotiable", "Sleep is part of the test strategy. Fatigue increases the chance of destructive migrations, missed authorisation checks and silent data loss. When energy drops, switch to low-risk documentation or test-writing; do not perform irreversible work.", fill="FDECEF", accent=RED)

doc.add_page_break()
heading(doc, "3. Starting position", 1)
paragraph(doc, "The current repository contains valuable domain knowledge but is not yet a safe migration source as a whole. Copy only code that has been classified and tested.")

baseline = doc.add_table(rows=1, cols=4)
baseline.style = "Table Grid"
table_geometry(baseline, [1900, 2250, 2600, 2610])
for i, text in enumerate(("Area", "Current condition", "Decision", "Reason")):
    shade(baseline.rows[0].cells[i], NAVY)
    font(baseline.rows[0].cells[i].paragraphs[0].add_run(text), size=8.8, bold=True, color=WHITE)
repeat_header(baseline.rows[0])
for values in [
    ("EF Core schema", "Broad and established", "Reuse", "Avoid unnecessary data migration"),
    ("Identity and privileges", "Good concept; integration gaps", "Reuse and harden", "Preserve users while fixing claims and scope"),
    ("Legacy Excel import", "Large, memory-backed workflow", "Retire after parity", "Not durable or sufficiently auditable"),
    ("Staged import", "Strong design; incomplete wiring", "Complete first", "Best first vertical slice"),
    ("Report calculations", "Substantial business logic", "Protect with tests", "High rewrite risk"),
    ("Razor/jQuery UI", "Tightly coupled and duplicated", "Replace incrementally", "Angular becomes the presentation layer"),
    ("Automated tests", "No usable test project", "Add immediately", "Required for fast, confident migration"),
]:
    row = baseline.add_row()
    for i, value in enumerate(values):
        font(row.cells[i].paragraphs[0].add_run(value), size=8.4, bold=(i in (0, 2)), color=NAVY if i in (0, 2) else INK)

heading(doc, "Known blockers to resolve first", 2)
for item in [
    "The existing solution currently fails to build because staged-import DbSet properties are missing and YearHistory is ambiguous.",
    "The staged-import services and options are not registered in dependency injection.",
    "The custom claims principal factory exists but is not registered.",
    "District and station claims are labels; query-level data scoping is not implemented.",
    "The staged-import migration exists, but repository inspection cannot prove whether it was applied.",
    "DevExtreme reports evaluation/licensing warnings; entitlement must be confirmed before deployment.",
    "A hardcoded bootstrap administrator password must not be carried into the new system.",
]:
    bullet(doc, item)

heading(doc, "4. Target architecture", 1)
arch = doc.add_table(rows=1, cols=2)
arch.style = "Table Grid"
table_geometry(arch, [2600, 6760])
for i, text in enumerate(("Layer", "Responsibility")):
    shade(arch.rows[0].cells[i], CYAN if i == 0 else NAVY)
    font(arch.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=NAVY if i == 0 else WHITE)
repeat_header(arch.rows[0])
for label, value in [
    ("Angular + DevExtreme", "Routes, pages, forms, grids, charts, loading/error states and permission-aware navigation."),
    ("ASP.NET Core API", "HTTP contracts, authentication, authorisation, antiforgery, ProblemDetails, uploads and downloads."),
    ("Core", "Business services, DTOs, import workflow, validation, reports, EF Core and document generation."),
    ("SQL Server", "Existing operational, summary, Identity, lookup and staged-import data."),
    ("Protected file storage", "Original workbooks, crash attachments and generated reports outside public web storage."),
]:
    row = arch.add_row()
    shade(row.cells[0], PALE_CYAN)
    font(row.cells[0].paragraphs[0].add_run(label), size=9, bold=True, color=NAVY)
    font(row.cells[1].paragraphs[0].add_run(value), size=9)

heading(doc, "Runtime solution", 2)
for item in [
    "AccidentReports.Api — API controllers, middleware, security configuration and application startup.",
    "AccidentReports.Core — entities, EF Core, feature services, validation, imports, reports and document generation.",
    "AccidentReports.Tests — unit, integration, API and regression tests.",
    "accident-reports-client — Angular and DevExtreme, created after the first API feature is stable.",
]:
    bullet(doc, item)

doc.add_page_break()
heading(doc, "5. Critical-path migration phases", 1)
paragraph(doc, "Effort is expressed in focused working days, not calendar promises. Each phase ends with an observable exit gate; if the gate fails, fix it before starting the next phase.")

phase(doc, 0, "Freeze the baseline and restore a green build", "1–2 days",
      "A trusted source baseline that builds and can be compared against the new solution.",
      [
          "Create a database backup or restored development copy and record the connection target.",
          "Decide which uncommitted import/report files are authoritative and commit them deliberately.",
          "Add the missing staged-import DbSet properties and resolve the YearHistory conflict in the original branch or a dedicated stabilisation branch.",
          "Run the existing application’s key workflows and record reference totals/screenshots.",
          "Prepare sanitised regression copies of representative original and standard-template workbooks.",
      ],
      "Original solution builds with zero errors; reference data and workbook expectations are recorded; no migration has been applied blindly.")

phase(doc, 1, "Create the new backend foundation", "1–2 days",
      "API, Core and Tests build together and can connect safely to a development database.",
      [
          "Add the API → Core project reference and Tests → API/Core references.",
          "Configure controllers, ProblemDetails, health checks, structured logging and environment-based configuration.",
          "Move entities, AppDbContext and migration history without renaming tables or keys.",
          "Register ASP.NET Identity, the custom claims principal factory and privilege policies.",
          "Create a single Core dependency-injection extension and one API health endpoint.",
          "Add smoke tests proving Core construction, database connectivity and API startup.",
      ],
      "dotnet build and dotnet test pass; /health responds; database schema comparison shows no accidental destructive change.")

phase(doc, 2, "Complete the staged-import backend", "3–4 days",
      "A workbook can be safely uploaded, staged, reviewed, approved and committed entirely through tested APIs.",
      [
          "Move the staged entities, options, repository, intake, detector, parser, validator, review and commit services into Core.",
          "Preserve original values, CASE numbers, demographic sections and workbook fingerprints.",
          "Add station-alias, Excel serial-time, duplicate and total-reconciliation rules discovered from real workbooks.",
          "Introduce granular privileges: view, upload, review, approve, commit and cancel.",
          "Create ImportBatches API endpoints with DTOs, pagination, cancellation and ProblemDetails errors.",
          "Test the original, converted and filled Ehlanzeni workbooks as regression cases.",
          "Retain the legacy importer only as a comparison oracle; do not let both pipelines write production data.",
      ],
      "The same representative workbook produces reconciled totals, preserves source evidence, flags known issues and commits exactly once.")

phase(doc, 3, "Create the Angular shell and import experience", "3–4 days",
      "The first Angular vertical slice is usable end to end against the new API.",
      [
          "Scaffold Angular with routing, strict TypeScript and SCSS; install licensed DevExtreme packages through npm.",
          "Create the vivid cyan/aqua/white/dark-blue design system, responsive shell and accessible navigation.",
          "Implement same-origin Identity cookie authentication, /api/auth/me, route guards and antiforgery handling.",
          "Build typed API services, one error interceptor, loading states and notification patterns.",
          "Implement batch history, workbook upload, staged row grid, issue details, review actions, approval and commit confirmation.",
          "Run keyboard, narrow-screen, forbidden-state and expired-session checks.",
      ],
      "A permitted user completes upload-to-commit in Angular; an under-privileged user is denied by the API; the MVC import remains available as fallback.")

doc.add_page_break()
phase(doc, 4, "Migrate crash registry and capture", "5–8 days",
      "Operational users can search, inspect and capture crash records without depending on Razor pages.",
      [
          "Extract crash search/detail queries from controllers into Core query services.",
          "Create paged CrashRecords endpoints and stable detail DTOs.",
          "Migrate quick capture first, then full capture, persons, vehicles, witnesses and attachments.",
          "Store attachments privately and serve them through authorised download endpoints.",
          "Enforce district/station data scope in every query and mutation.",
          "Compare new records with the old application’s expected database results.",
      ],
      "Authorised users can complete the agreed capture workflows; scope tests pass; no attachment is publicly addressable.")

phase(doc, 5, "Consolidate reports, dashboard and insights", "4–6 days",
      "All official reporting and analytical views use shared, tested backend calculations.",
      [
          "Consolidate monthly, quarterly, six-month and nine-month periods behind a shared resolver and data service.",
          "Keep standby and five-year specialisations where their data shapes genuinely differ.",
          "Wrap existing Word generation behind one document-generator contract; keep server-side generation.",
          "Add report history, status and authorised download endpoints.",
          "Build Angular Reports, Dashboard and Insights pages from API DTOs.",
          "Regression-test crash, fatality, serious, slight, route and demographic totals for known periods.",
      ],
      "Generated documents and preview totals match approved reference outputs; reports are authorised, auditable and downloadable.")

phase(doc, 6, "Administration, hardening and cutover", "2–4 days",
      "The new application is supportable, secure and ready for controlled acceptance.",
      [
          "Migrate lookups, users, roles and privilege administration.",
          "Replace text-based user station/district scope with reconciled lookup IDs where approved.",
          "Run dependency, licence, security, accessibility, performance and backup/restore checks.",
          "Prepare deployment, rollback, migration and support runbooks.",
          "Conduct user acceptance by feature and record formal gaps.",
          "Retire each MVC feature only after its replacement is accepted; archive the old application after full cutover.",
      ],
      "Acceptance evidence is signed off, rollback is tested, monitoring is active and the old system is no longer required for accepted workflows.")

heading(doc, "6. Fast-track release boundaries", 1)
tracks = doc.add_table(rows=1, cols=4)
tracks.style = "Table Grid"
table_geometry(tracks, [1750, 2500, 2550, 2560])
for i, text in enumerate(("Release", "Includes", "Excludes", "Ready when")):
    shade(tracks.rows[0].cells[i], NAVY)
    font(tracks.rows[0].cells[i].paragraphs[0].add_run(text), size=8.8, bold=True, color=WHITE)
repeat_header(tracks.rows[0])
for values in [
    ("MVP", "Foundation, auth, staged imports and quality review", "Full crash capture, reports and administration", "Upload-to-commit is secure and tested"),
    ("Operational", "Crash registry, quick/full capture and protected attachments", "Optional visual polish", "Operational users accept core workflows"),
    ("Complete", "Reports, dashboard, insights, lookups, users and roles", "Future enhancements", "All agreed MVC features have accepted replacements"),
]:
    row = tracks.add_row()
    for i, value in enumerate(values):
        font(row.cells[i].paragraphs[0].add_run(value), size=8.6, bold=(i == 0), color=CYAN if i == 0 else INK)

callout(doc, "Fastest sensible target", "Ship the import-focused MVP first. It proves the API, authentication, permissions, database migration, DevExtreme grids, file handling and the hardest data-quality workflow in one slice.")

doc.add_page_break()
heading(doc, "7. Database migration strategy", 1)
paragraph(doc, "Angular does not require a new database. Database changes should be additive, reviewed and independently reversible wherever practical.")
for item in [
    "Keep the existing SQL Server database, table names, primary keys and Identity records.",
    "Designate AccidentReports.Core as the sole EF migration owner after cutover.",
    "Never allow the MVC project and new API to generate competing migrations against the same database.",
    "Apply migrations first to a restored development database, then test rollback or restore procedures.",
    "Separate schema deployment from application startup; do not auto-migrate the production database on every API launch.",
    "Back up before every production migration and record the exact application version paired with it.",
]:
    bullet(doc, item)

heading(doc, "Required or recommended schema work", 2)
db = doc.add_table(rows=1, cols=3)
db.style = "Table Grid"
table_geometry(db, [2850, 3500, 3010])
for i, text in enumerate(("Change", "Purpose", "Timing")):
    shade(db.rows[0].cells[i], NAVY)
    font(db.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
repeat_header(db.rows[0])
for change, purpose, timing in [
    ("Complete staged-import tables and mappings", "Durable upload, review and commit", "Before import API"),
    ("Reviewer/resolver user foreign keys", "Stronger audit integrity", "With staged pipeline review"),
    ("Station alias table", "Controlled spelling normalisation", "Import MVP or shortly after"),
    ("Generated report history", "Audited generation and downloads", "Before Reports UI"),
    ("Attachment metadata and private storage", "Authorised file access", "Before capture migration"),
    ("User station_id and district_id", "Enforceable operational scope", "After lookup reconciliation"),
    ("RowVersion concurrency columns", "Prevent silent review/edit overwrites", "Before multi-user acceptance"),
]:
    row = db.add_row()
    for i, value in enumerate((change, purpose, timing)):
        font(row.cells[i].paragraphs[0].add_run(value), size=8.7, bold=(i == 0), color=NAVY if i == 0 else INK)

heading(doc, "Migration safety checklist", 2)
for item in [
    "Migration script reviewed; no unexpected drop, rename or data conversion.",
    "Backup/restore procedure tested on a development copy.",
    "Existing row counts and key report totals captured before migration.",
    "Migration applied to development and staging before production.",
    "Application smoke tests and reconciliation queries pass afterward.",
    "Rollback decision point and responsible person recorded.",
]:
    bullet(doc, item, checked=False)

heading(doc, "8. Security and quality gates", 1)
quality = doc.add_table(rows=1, cols=3)
quality.style = "Table Grid"
table_geometry(quality, [2300, 3800, 3260])
for i, text in enumerate(("Gate", "Evidence", "Failure means")):
    shade(quality.rows[0].cells[i], NAVY)
    font(quality.rows[0].cells[i].paragraphs[0].add_run(text), size=9, bold=True, color=WHITE)
repeat_header(quality.rows[0])
for gate, evidence, failure in [
    ("Build", "Zero compile errors", "Stop; do not layer new work on a broken baseline"),
    ("Tests", "Relevant unit/integration tests pass", "Fix or explicitly quarantine the failing change"),
    ("Data", "Counts and report totals reconcile", "Do not commit or deploy"),
    ("Authorisation", "Unauthorised and wrong-scope tests fail correctly", "Release is blocked"),
    ("Files", "Type, size, path and download permission checks pass", "Release is blocked"),
    ("UX", "Loading, empty, error, forbidden and narrow-screen states work", "Feature remains in preview"),
    ("Rollback", "Previous application/database state can be restored", "Production migration is blocked"),
]:
    row = quality.add_row()
    for i, value in enumerate((gate, evidence, failure)):
        font(row.cells[i].paragraphs[0].add_run(value), size=8.6, bold=(i == 0), color=RED if i == 2 else (NAVY if i == 0 else INK))

heading(doc, "9. Solo-developer execution rhythm", 1)
paragraph(doc, "Use a rhythm that protects momentum and makes every day resumable. The goal is sustained velocity, not heroic context loss.")
for item in [
    "Start: choose one measurable outcome and its acceptance test.",
    "Deep-work block 1: implement the smallest backend or frontend change that proves progress.",
    "Midpoint: build and run focused tests before expanding scope.",
    "Deep-work block 2: complete the vertical path or fix the discovered gap.",
    "Close: run build/tests, write a short decision note, commit a coherent change and record tomorrow’s first action.",
    "Fatigue rule: no production credentials, destructive migrations, privilege changes or irreversible commits when exhausted.",
]:
    bullet(doc, item)

heading(doc, "10. Migration board", 1)
board = doc.add_table(rows=1, cols=5)
board.style = "Table Grid"
table_geometry(board, [700, 2050, 3000, 1700, 1910])
for i, text in enumerate(("Pri", "Slice", "Definition", "Dependency", "Status")):
    shade(board.rows[0].cells[i], NAVY)
    font(board.rows[0].cells[i].paragraphs[0].add_run(text), size=8.6, bold=True, color=WHITE)
repeat_header(board.rows[0])
for values in [
    ("P0", "Stabilise baseline", "Build green; DB state known", "None", "Not started"),
    ("P0", "Backend foundation", "API/Core/Tests operational", "Baseline", "Not started"),
    ("P0", "Import backend", "Upload-to-commit API works", "Foundation", "Not started"),
    ("P0", "Angular import slice", "Users complete review workflow", "Import API", "Not started"),
    ("P1", "Crash registry", "Paged search and details", "Auth/scope", "Backlog"),
    ("P1", "Crash capture", "Quick and full capture", "Registry", "Backlog"),
    ("P1", "Reports", "Shared periods and downloads", "Stable data", "Backlog"),
    ("P2", "Dashboard/insights", "API-backed operational intelligence", "Reports", "Backlog"),
    ("P2", "Administration", "Users, roles, privileges, lookups", "Auth/scope", "Backlog"),
]:
    row = board.add_row()
    for i, value in enumerate(values):
        colour = RED if value[0] == "P0" else (GOLD if value[0] == "P1" else MUTED)
        font(row.cells[i].paragraphs[0].add_run(value), size=8.3, bold=(i in (0, 1)), color=colour if i == 0 else (NAVY if i == 1 else INK))

doc.add_page_break()
heading(doc, "11. Definition of done", 1)
for item in [
    "The feature builds from documented commands and all relevant automated tests pass.",
    "The API uses DTOs, ProblemDetails, cancellation and explicit authorisation policies.",
    "District/station scope is enforced by server-side queries where applicable.",
    "Database changes are reviewed, backed up and reconciled.",
    "Angular handles loading, empty, success, validation, forbidden and unexpected-error states.",
    "Keyboard flow, labels, contrast and narrow-screen layout are usable.",
    "No original workbook value is silently overwritten; corrections retain original and suggested values.",
    "Operational and security-relevant actions are auditable.",
    "Deployment and rollback steps are documented.",
    "The MVC feature remains available until acceptance is recorded.",
]:
    bullet(doc, item, checked=False)

heading(doc, "12. Decisions to record before work begins", 1)
for item in [
    "Exact development and production database targets",
    "Whether the staged-import migration has already been applied anywhere",
    "DevExtreme/DevExpress licence owner and approved version",
    "Same-origin hosting decision for Angular and the API",
    "Who may upload, review, approve and commit imports",
    "Approved station/district master data and alias ownership",
    "Approved reference report documents and reconciliation periods",
    "Deployment owner, backup owner and rollback authority",
]:
    bullet(doc, item, checked=False)

callout(doc, "First move tomorrow", "Do not scaffold Angular first. Restore the original repository to a green build, confirm the development database state, create the Tests project and prove the new API can start. Then migrate the staged import backend as the first complete slice.", fill=PALE_CYAN, accent=CYAN)
paragraph(doc, "You are not starting over. You are extracting the strongest parts of the system into a safer shape, proving them one slice at a time, and replacing only what has earned replacement.", size=12, bold=True, color=NAVY, before=8, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)

doc.core_properties.title = "Accident Reports Admin System Rapid Migration Strategy"
doc.core_properties.subject = "Backend-led migration from ASP.NET Core MVC to Angular, ASP.NET Core API and DevExtreme"
doc.core_properties.author = "Accident Reports Admin System Project"
doc.core_properties.keywords = "Angular, ASP.NET Core API, DevExtreme, migration, data cleaning, MPRIMS"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
